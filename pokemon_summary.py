#!/usr/bin/env python3
"""
Pokemon Data Summary Generator
Creates a compact summary of all Pokemon data in different formats
"""

import os
import re
import sys
import csv
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor

# Import the main parser
from pokemon_parser import PokemonParser

def create_summary_document(parser, output_file="pokemon_summary.docx"):
    """Create a compact summary document"""
    doc = Document()
    
    # Add title
    title = doc.add_heading('Pokemon ROM Hack - Quick Reference', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add stats
    doc.add_paragraph(f"Total Pokemon: {len(parser.species_data)}")
    doc.add_paragraph("")
    
    # Create a table with all Pokemon data
    # Headers: Name, Types, BST, HP, Att, Def, SpA, SpD, Spe, Abilities
    table = doc.add_table(rows=1, cols=10)
    table.style = 'Table Grid'
    
    # Set header row
    headers = ['Pokemon', 'Type(s)', 'BST', 'HP', 'Att', 'Def', 'SpA', 'SpD', 'Spe', 'Abilities']
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        # Make header bold
        run = cell.paragraphs[0].runs[0]
        run.font.bold = True
    
    # Sort Pokemon by display name for better readability
    sorted_pokemon = sorted(parser.species_data.values(), key=lambda x: x['display_name'])
    
    # Add data rows
    for pokemon in sorted_pokemon:
        row = table.add_row()
        
        # Get the appropriate display name with form information
        form_display_name = parser.get_form_display_name(pokemon['name'], pokemon['display_name'])
        
        # Pokemon name
        row.cells[0].text = form_display_name
        
        # Types
        types_str = " / ".join(pokemon['types']) if pokemon['types'] else "?"
        row.cells[1].text = types_str
        
        # Calculate BST and individual stats
        stats = pokemon['stats']
        bst = sum(stats.values()) if stats else 0
        row.cells[2].text = str(bst)
        
        # Individual stats
        stat_order = ['HP', 'Attack', 'Defense', 'Sp. Attack', 'Sp. Defense', 'Speed']
        for i, stat_name in enumerate(stat_order):
            value = stats.get(stat_name, 0)
            row.cells[3 + i].text = str(value)
        
        # Abilities
        abilities_str = ", ".join(pokemon['abilities']) if pokemon['abilities'] else "?"
        row.cells[9].text = abilities_str
    
    # Save document with error handling
    output_path = parser.base_path / output_file
    
    # Try to save, if permission denied, try with a different name
    attempt = 0
    max_attempts = 5
    while attempt < max_attempts:
        try:
            if attempt == 0:
                save_path = output_path
            else:
                # Add a number to the filename
                stem = output_path.stem
                suffix = output_path.suffix
                save_path = output_path.parent / f"{stem}_{attempt}{suffix}"
            
            doc.save(save_path)
            print(f"Summary document saved to: {save_path}")
            return save_path
            
        except PermissionError as e:
            attempt += 1
            if attempt == 1:
                print(f"Permission denied for {output_path}")
                print("The file might be open in Word. Trying alternative filename...")
            elif attempt < max_attempts:
                print(f"Attempt {attempt} failed, trying {save_path}...")
            else:
                print(f"Failed to save after {max_attempts} attempts.")
                print("Please close the Word document if it's open and try again.")
                raise e
        except Exception as e:
            print(f"Error saving document: {e}")
            raise e

def create_csv_export(parser, output_file="pokemon_data.csv"):
    """Export Pokemon data to CSV format"""
    output_path = parser.base_path / output_file
    
    # Try to save CSV with error handling
    attempt = 0
    max_attempts = 5
    while attempt < max_attempts:
        try:
            if attempt == 0:
                save_path = output_path
            else:
                # Add a number to the filename
                stem = output_path.stem
                suffix = output_path.suffix
                save_path = output_path.parent / f"{stem}_{attempt}{suffix}"
            
            with open(save_path, 'w', newline='', encoding='utf-8') as csvfile:
                fieldnames = [
                    'Name', 'Display_Name', 'Type_1', 'Type_2', 'BST',
                    'HP', 'Attack', 'Defense', 'Sp_Attack', 'Sp_Defense', 'Speed',
                    'Ability_1', 'Ability_2', 'Ability_3'
                ]
                writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
                
                writer.writeheader()
                
                # Sort Pokemon by display name instead of internal name
                sorted_pokemon = sorted(parser.species_data.values(), key=lambda x: x['display_name'])
                
                for pokemon in sorted_pokemon:
                    stats = pokemon['stats']
                    types = pokemon['types'] + [''] * (2 - len(pokemon['types']))  # Pad to 2 types
                    abilities = pokemon['abilities'] + [''] * (3 - len(pokemon['abilities']))  # Pad to 3 abilities
                    
                    # Get the appropriate display name with form information
                    form_display_name = parser.get_form_display_name(pokemon['name'], pokemon['display_name'])
                    
                    row = {
                        'Name': pokemon['name'],
                        'Display_Name': form_display_name,
                        'Type_1': types[0] if len(types) > 0 else '',
                        'Type_2': types[1] if len(types) > 1 else '',
                        'BST': sum(stats.values()) if stats else 0,
                        'HP': stats.get('HP', 0),
                        'Attack': stats.get('Attack', 0),
                        'Defense': stats.get('Defense', 0),
                        'Sp_Attack': stats.get('Sp. Attack', 0),
                        'Sp_Defense': stats.get('Sp. Defense', 0),
                        'Speed': stats.get('Speed', 0),
                        'Ability_1': abilities[0] if len(abilities) > 0 else '',
                        'Ability_2': abilities[1] if len(abilities) > 1 else '',
                        'Ability_3': abilities[2] if len(abilities) > 2 else ''
                    }
                    writer.writerow(row)
            
            print(f"CSV export saved to: {save_path}")
            return save_path
            
        except PermissionError as e:
            attempt += 1
            if attempt == 1:
                print(f"Permission denied for {output_path}")
                print("The CSV file might be open in Excel. Trying alternative filename...")
            elif attempt < max_attempts:
                print(f"Attempt {attempt} failed, trying {save_path}...")
            else:
                print(f"Failed to save CSV after {max_attempts} attempts.")
                print("Please close the CSV file if it's open and try again.")
                raise e
        except Exception as e:
            print(f"Error saving CSV: {e}")
            raise e

def main():
    # Get the current directory
    base_path = Path(__file__).parent
    
    print("Pokemon ROM Hack Data Summary Generator")
    print("=" * 50)
    
    # Initialize parser and parse data
    parser = PokemonParser(base_path)
    parser.load_constants()
    parser.parse_species_files()
    
    print(f"Found {len(parser.species_data)} Pokemon with data")
    
    # Create summary document
    print("Creating summary document...")
    create_summary_document(parser)
    
    # Create CSV export
    print("Creating CSV export...")
    create_csv_export(parser)
    
    print("All exports complete!")

if __name__ == "__main__":
    main()
